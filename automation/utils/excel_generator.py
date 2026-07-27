"""
Excel Report Generator – creates all 4 required Excel workbooks.
Uses openpyxl for .xlsx generation.
"""

import os
import time
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from config.settings import REPORTS_DIR, BASE_URL

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


# ── Style Constants ───────────────────────────────────────────────────────
if HAS_OPENPYXL:
    HEADER_FONT = Font(name="Calibri", bold=True, size=12, color="FFFFFF")
    HEADER_FILL = PatternFill(start_color="1A1A3E", end_color="1A1A3E", fill_type="solid")
    PASS_FILL = PatternFill(start_color="C8E6C9", end_color="C8E6C9", fill_type="solid")
    FAIL_FILL = PatternFill(start_color="FFCDD2", end_color="FFCDD2", fill_type="solid")
    SKIP_FILL = PatternFill(start_color="FFF9C4", end_color="FFF9C4", fill_type="solid")
    PASS_FONT = Font(name="Calibri", color="2E7D32", bold=True)
    FAIL_FONT = Font(name="Calibri", color="C62828", bold=True)
    SKIP_FONT = Font(name="Calibri", color="F57F17", bold=True)
    TITLE_FONT = Font(name="Calibri", bold=True, size=16, color="1A237E")
    BORDER = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    CENTER = Alignment(horizontal="center", vertical="center")


def _style_header_row(ws, columns):
    """Apply header styling to the first row."""
    if not HAS_OPENPYXL:
        return
    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER


def _auto_width(ws):
    """Auto-adjust column widths."""
    if not HAS_OPENPYXL:
        return
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 4, 60)


def _add_result_row(ws, row_idx, result):
    """Add a test result row with status-based styling."""
    if not HAS_OPENPYXL:
        return
    values = [
        result.get("test_id", ""),
        result.get("module", ""),
        result.get("test_name", ""),
        result.get("status", ""),
        f"{result.get('duration', 0):.2f}s",
        result.get("priority", "Medium"),
        result.get("error", "")[:200] if result.get("error") else "",
        result.get("precondition", "App loaded"),
        result.get("expected", ""),
        result.get("actual", ""),
    ]
    for col_idx, val in enumerate(values, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.border = BORDER
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        if col_idx == 4:  # Status column
            if val == "PASSED":
                cell.fill = PASS_FILL
                cell.font = PASS_FONT
            elif val == "FAILED":
                cell.fill = FAIL_FILL
                cell.font = FAIL_FONT
            elif val == "SKIPPED":
                cell.fill = SKIP_FILL
                cell.font = SKIP_FONT


def generate_automation_test_report(results: list, output_dir: str = None):
    """Generate the main Automation_Test_Report.xlsx with 6 sheets."""
    if not HAS_OPENPYXL:
        print("[EXCEL] openpyxl not available, skipping Excel generation")
        return None

    if output_dir is None:
        output_dir = REPORTS_DIR
    os.makedirs(output_dir, exist_ok=True)

    wb = Workbook()
    columns = ["Test ID", "Module", "Test Name", "Status", "Execution Time",
               "Priority", "Error Details", "Precondition", "Expected Result", "Actual Result"]

    # Sheet 1: All Executed Tests
    ws1 = wb.active
    ws1.title = "Executed Test Cases"
    _style_header_row(ws1, columns)
    for idx, r in enumerate(results, 2):
        _add_result_row(ws1, idx, r)
    _auto_width(ws1)

    # Sheet 2: Passed Tests
    ws2 = wb.create_sheet("Passed Tests")
    passed = [r for r in results if r["status"] == "PASSED"]
    _style_header_row(ws2, columns)
    for idx, r in enumerate(passed, 2):
        _add_result_row(ws2, idx, r)
    _auto_width(ws2)

    # Sheet 3: Failed Tests
    ws3 = wb.create_sheet("Failed Tests")
    failed = [r for r in results if r["status"] == "FAILED"]
    _style_header_row(ws3, columns)
    for idx, r in enumerate(failed, 2):
        _add_result_row(ws3, idx, r)
    _auto_width(ws3)

    # Sheet 4: Skipped Tests
    ws4 = wb.create_sheet("Skipped Tests")
    skipped = [r for r in results if r["status"] == "SKIPPED"]
    _style_header_row(ws4, columns)
    for idx, r in enumerate(skipped, 2):
        _add_result_row(ws4, idx, r)
    _auto_width(ws4)

    # Sheet 5: Execution Metrics
    ws5 = wb.create_sheet("Execution Metrics")
    metrics = [
        ["Metric", "Value"],
        ["Total Test Cases", len(results)],
        ["Passed", len(passed)],
        ["Failed", len(failed)],
        ["Skipped", len(skipped)],
        ["Pass Rate (%)", f"{len(passed)/len(results)*100:.1f}" if results else "0"],
        ["Total Duration (s)", f"{sum(r.get('duration', 0) for r in results):.1f}"],
        ["Avg Duration (s)", f"{sum(r.get('duration', 0) for r in results)/len(results):.2f}" if results else "0"],
        ["Deployment URL", BASE_URL],
        ["Execution Date", time.strftime("%Y-%m-%d %H:%M:%S")],
    ]
    for row_idx, row_data in enumerate(metrics, 1):
        for col_idx, val in enumerate(row_data, 1):
            cell = ws5.cell(row=row_idx, column=col_idx, value=val)
            cell.border = BORDER
            if row_idx == 1:
                cell.font = HEADER_FONT
                cell.fill = HEADER_FILL
    _auto_width(ws5)

    # Sheet 6: Defect Summary
    ws6 = wb.create_sheet("Defect Summary")
    defect_cols = ["Defect ID", "Module", "Test Name", "Severity", "Description", "Steps to Reproduce"]
    _style_header_row(ws6, defect_cols)
    for idx, r in enumerate(failed, 2):
        ws6.cell(row=idx, column=1, value=f"DEF-{idx-1:04d}").border = BORDER
        ws6.cell(row=idx, column=2, value=r.get("module", "")).border = BORDER
        ws6.cell(row=idx, column=3, value=r.get("test_name", "")).border = BORDER
        ws6.cell(row=idx, column=4, value=r.get("priority", "Medium")).border = BORDER
        ws6.cell(row=idx, column=5, value=r.get("error", "")[:200]).border = BORDER
        ws6.cell(row=idx, column=6, value=f"1. Navigate to {BASE_URL}\n2. Execute {r.get('test_name', '')}").border = BORDER
    _auto_width(ws6)

    filepath = os.path.join(output_dir, "Automation_Test_Report.xlsx")
    wb.save(filepath)
    return filepath


def generate_failed_report(results: list, output_dir: str = None):
    """Generate Failed_Test_Cases.xlsx."""
    if not HAS_OPENPYXL:
        return None
    if output_dir is None:
        output_dir = REPORTS_DIR

    wb = Workbook()
    ws = wb.active
    ws.title = "Failed Test Cases"
    cols = ["Test ID", "Module", "Test Name", "Priority", "Error", "Stack Trace", "Screenshot", "Duration"]
    _style_header_row(ws, cols)

    failed = [r for r in results if r["status"] == "FAILED"]
    for idx, r in enumerate(failed, 2):
        ws.cell(row=idx, column=1, value=r.get("test_id", "")).border = BORDER
        ws.cell(row=idx, column=2, value=r.get("module", "")).border = BORDER
        ws.cell(row=idx, column=3, value=r.get("test_name", "")).border = BORDER
        ws.cell(row=idx, column=4, value=r.get("priority", "Medium")).border = BORDER
        ws.cell(row=idx, column=5, value=r.get("error", "")[:500]).border = BORDER
        ws.cell(row=idx, column=6, value=r.get("stack_trace", "")[:500]).border = BORDER
        ws.cell(row=idx, column=7, value=r.get("screenshot", "")).border = BORDER
        ws.cell(row=idx, column=8, value=f"{r.get('duration', 0):.2f}s").border = BORDER
    _auto_width(ws)

    filepath = os.path.join(output_dir, "Failed_Test_Cases.xlsx")
    wb.save(filepath)
    return filepath


def generate_passed_report(results: list, output_dir: str = None):
    """Generate Passed_Test_Cases.xlsx."""
    if not HAS_OPENPYXL:
        return None
    if output_dir is None:
        output_dir = REPORTS_DIR

    wb = Workbook()
    ws = wb.active
    ws.title = "Passed Test Cases"
    cols = ["Test ID", "Module", "Test Name", "Priority", "Duration", "Verified At"]
    _style_header_row(ws, cols)

    passed = [r for r in results if r["status"] == "PASSED"]
    for idx, r in enumerate(passed, 2):
        ws.cell(row=idx, column=1, value=r.get("test_id", "")).border = BORDER
        ws.cell(row=idx, column=2, value=r.get("module", "")).border = BORDER
        ws.cell(row=idx, column=3, value=r.get("test_name", "")).border = BORDER
        ws.cell(row=idx, column=4, value=r.get("priority", "Medium")).border = BORDER
        ws.cell(row=idx, column=5, value=f"{r.get('duration', 0):.2f}s").border = BORDER
        ws.cell(row=idx, column=6, value=time.strftime("%Y-%m-%d %H:%M:%S")).border = BORDER
    _auto_width(ws)

    filepath = os.path.join(output_dir, "Passed_Test_Cases.xlsx")
    wb.save(filepath)
    return filepath


def generate_summary_report(results: list, output_dir: str = None):
    """Generate Summary_Report.xlsx."""
    if not HAS_OPENPYXL:
        return None
    if output_dir is None:
        output_dir = REPORTS_DIR

    wb = Workbook()
    ws = wb.active
    ws.title = "Executive Summary"

    passed = len([r for r in results if r["status"] == "PASSED"])
    failed = len([r for r in results if r["status"] == "FAILED"])
    skipped = len([r for r in results if r["status"] == "SKIPPED"])
    total = len(results)

    summary_data = [
        ["MediRoute E2E Test Execution Summary"],
        [""],
        ["Metric", "Value"],
        ["Deployment URL", BASE_URL],
        ["Execution Date", time.strftime("%Y-%m-%d %H:%M:%S")],
        ["Total Test Cases", total],
        ["Passed", passed],
        ["Failed", failed],
        ["Skipped", skipped],
        ["Pass Rate", f"{passed/total*100:.1f}%" if total else "0%"],
        ["Total Duration", f"{sum(r.get('duration', 0) for r in results):.1f}s"],
        [""],
        ["Module Breakdown"],
    ]

    for row_idx, row in enumerate(summary_data, 1):
        for col_idx, val in enumerate(row, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = BORDER

    # Title formatting
    ws.cell(row=1, column=1).font = TITLE_FONT
    ws.merge_cells("A1:B1")

    # Module breakdown
    modules = {}
    for r in results:
        m = r["module"]
        if m not in modules:
            modules[m] = {"passed": 0, "failed": 0, "skipped": 0, "total": 0}
        modules[m]["total"] += 1
        modules[m][r["status"].lower()] = modules[m].get(r["status"].lower(), 0) + 1

    row_start = len(summary_data) + 1
    ws.cell(row=row_start, column=1, value="Module").font = HEADER_FONT
    ws.cell(row=row_start, column=1).fill = HEADER_FILL
    ws.cell(row=row_start, column=2, value="Passed").font = HEADER_FONT
    ws.cell(row=row_start, column=2).fill = HEADER_FILL
    ws.cell(row=row_start, column=3, value="Failed").font = HEADER_FONT
    ws.cell(row=row_start, column=3).fill = HEADER_FILL
    ws.cell(row=row_start, column=4, value="Rate").font = HEADER_FONT
    ws.cell(row=row_start, column=4).fill = HEADER_FILL

    for idx, (m, stats) in enumerate(sorted(modules.items()), row_start + 1):
        rate = stats["passed"] / stats["total"] * 100 if stats["total"] else 0
        ws.cell(row=idx, column=1, value=m).border = BORDER
        ws.cell(row=idx, column=2, value=stats["passed"]).border = BORDER
        ws.cell(row=idx, column=3, value=stats["failed"]).border = BORDER
        ws.cell(row=idx, column=4, value=f"{rate:.1f}%").border = BORDER

    _auto_width(ws)

    filepath = os.path.join(output_dir, "Summary_Report.xlsx")
    wb.save(filepath)
    return filepath


def generate_all_excel_reports(results: list, output_dir: str = None):
    """Generate all 4 Excel workbooks."""
    if output_dir is None:
        output_dir = REPORTS_DIR
    os.makedirs(output_dir, exist_ok=True)

    files = []
    files.append(generate_automation_test_report(results, output_dir))
    files.append(generate_failed_report(results, output_dir))
    files.append(generate_passed_report(results, output_dir))
    files.append(generate_summary_report(results, output_dir))

    return [f for f in files if f]
