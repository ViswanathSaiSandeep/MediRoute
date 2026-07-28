"""
Category-specific Excel and JSON report generator.
Produces separate xlsx per test category for CI artifacts.
"""

import json
import os
import sys
import time

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from config.settings import BASE_URL

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

CATEGORY_FILES = {
    "Selenium Website Tests": "Selenium_Website_Tests_Report.xlsx",
    "Appium Android Tests": "Appium_Android_Tests_Report.xlsx",
    "Unit Tests API": "Unit_Tests_API_Report.xlsx",
    "Validation Tests": "Validation_Tests_Report.xlsx",
    "Deployment Status": "Deployment_Status_Report.xlsx",
    "Load Testing Performance": "Load_Testing_Performance_Report.xlsx",
}

if HAS_OPENPYXL:
    HEADER_FONT = Font(name="Calibri", bold=True, size=12, color="FFFFFF")
    HEADER_FILL = PatternFill(start_color="1A237E", end_color="1A237E", fill_type="solid")
    PASS_FILL = PatternFill(start_color="C8E6C9", end_color="C8E6C9", fill_type="solid")
    FAIL_FILL = PatternFill(start_color="FFCDD2", end_color="FFCDD2", fill_type="solid")
    SKIP_FILL = PatternFill(start_color="FFF9C4", end_color="FFF9C4", fill_type="solid")
    BORDER = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )


def _style_header(ws, columns):
    if not HAS_OPENPYXL:
        return
    for i, col in enumerate(columns, 1):
        cell = ws.cell(row=1, column=i, value=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
        cell.border = BORDER


def _auto_width(ws):
    if not HAS_OPENPYXL:
        return
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 60)


def generate_category_reports(results: list, category: str, output_dir: str) -> dict:
    """Generate Excel + JSON reports for a test category."""
    os.makedirs(output_dir, exist_ok=True)
    passed = [r for r in results if r["status"] == "PASSED"]
    failed = [r for r in results if r["status"] == "FAILED"]
    skipped = [r for r in results if r["status"] == "SKIPPED"]
    total = len(results)

    summary = {
        "category": category,
        "total": total,
        "passed": len(passed),
        "failed": len(failed),
        "skipped": len(skipped),
        "pass_rate": round(len(passed) / total * 100, 2) if total else 0,
        "duration_seconds": round(sum(r.get("duration", 0) for r in results), 2),
        "deployment_url": BASE_URL,
        "execution_date": time.strftime("%Y-%m-%dT%H:%M:%SZ"),
    }

    # JSON results
    json_path = os.path.join(output_dir, "execution-results.json")
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump({"summary": summary, "results": results}, f, indent=2, default=str)

    # Summary markdown
    md_path = os.path.join(output_dir, "summary.md")
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(f"# {category}\n\n")
        f.write(f"| Metric | Value |\n|--------|-------|\n")
        f.write(f"| Total | {total} |\n| Passed | {len(passed)} |\n")
        f.write(f"| Failed | {len(failed)} |\n| Skipped | {len(skipped)} |\n")
        f.write(f"| Pass Rate | {summary['pass_rate']}% |\n")
        f.write(f"| URL | {BASE_URL} |\n")

    excel_path = None
    if HAS_OPENPYXL:
        wb = Workbook()
        cols = ["Test ID", "Module", "Test Name", "Status", "Execution Time",
                "Priority", "Precondition", "Expected Result", "Actual Result", "Error"]

        # Sheet 1: All tests
        ws1 = wb.active
        ws1.title = "Executed Test Cases"
        _style_header(ws1, cols)
        for i, r in enumerate(results, 2):
            row = [
                r.get("test_id"), r.get("module"), r.get("test_name"),
                r.get("status"), f"{r.get('duration', 0):.2f}s",
                r.get("priority"), r.get("precondition"),
                r.get("expected"), r.get("actual"), r.get("error", "")[:200],
            ]
            for j, val in enumerate(row, 1):
                cell = ws1.cell(row=i, column=j, value=val)
                cell.border = BORDER
                if j == 4:
                    if val == "PASSED":
                        cell.fill = PASS_FILL
                    elif val == "FAILED":
                        cell.fill = FAIL_FILL
                    elif val == "SKIPPED":
                        cell.fill = SKIP_FILL
        _auto_width(ws1)

        # Sheet 2-4: Passed/Failed/Skipped
        for sheet_name, subset in [("Passed Tests", passed), ("Failed Tests", failed), ("Skipped Tests", skipped)]:
            ws = wb.create_sheet(sheet_name)
            _style_header(ws, cols)
            for i, r in enumerate(subset, 2):
                row = [
                    r.get("test_id"), r.get("module"), r.get("test_name"),
                    r.get("status"), f"{r.get('duration', 0):.2f}s",
                    r.get("priority"), r.get("precondition"),
                    r.get("expected"), r.get("actual"), r.get("error", "")[:200],
                ]
                for j, val in enumerate(row, 1):
                    ws.cell(row=i, column=j, value=val).border = BORDER
            _auto_width(ws)

        # Sheet 5: Metrics
        ws5 = wb.create_sheet("Execution Metrics")
        metrics = [
            ["Metric", "Value"],
            ["Category", category],
            ["Total Test Cases", total],
            ["Passed", len(passed)],
            ["Failed", len(failed)],
            ["Skipped", len(skipped)],
            ["Pass Rate (%)", f"{summary['pass_rate']}"],
            ["Total Duration (s)", f"{summary['duration_seconds']}"],
            ["Deployment URL", BASE_URL],
            ["Execution Date", time.strftime("%Y-%m-%d %H:%M:%S")],
        ]
        for i, row in enumerate(metrics, 1):
            for j, val in enumerate(row, 1):
                cell = ws5.cell(row=i, column=j, value=val)
                cell.border = BORDER
                if i == 1:
                    cell.font = HEADER_FONT
                    cell.fill = HEADER_FILL
        _auto_width(ws5)

        # Sheet 6: Defect Summary
        ws6 = wb.create_sheet("Defect Summary")
        defect_cols = ["Defect ID", "Module", "Test Name", "Priority", "Description"]
        _style_header(ws6, defect_cols)
        for i, r in enumerate(failed, 2):
            ws6.cell(row=i, column=1, value=f"DEF-{i-1:04d}").border = BORDER
            ws6.cell(row=i, column=2, value=r.get("module")).border = BORDER
            ws6.cell(row=i, column=3, value=r.get("test_name")).border = BORDER
            ws6.cell(row=i, column=4, value=r.get("priority")).border = BORDER
            ws6.cell(row=i, column=5, value=r.get("error", "")[:300]).border = BORDER
        _auto_width(ws6)

        filename = CATEGORY_FILES.get(category, f"{category.replace(' ', '_')}_Report.xlsx")
        excel_path = os.path.join(output_dir, filename)
        wb.save(excel_path)

        # Also save standard names for compatibility
        for alt_name in ["Automation_Test_Report.xlsx", "Summary_Report.xlsx"]:
            alt_path = os.path.join(output_dir, alt_name)
            if alt_path != excel_path:
                wb.save(alt_path)

    return {"summary": summary, "excel": excel_path, "json": json_path, "markdown": md_path}
